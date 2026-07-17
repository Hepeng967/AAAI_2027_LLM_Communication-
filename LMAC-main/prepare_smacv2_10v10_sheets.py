#!/usr/bin/env python3
"""Add accurate SMACv2 10v10 observation/state descriptions to sc2.xlsx."""

from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
BOOK = ROOT / "src" / "llm_source" / "sc2.xlsx"


def row(name, low=0, high=1, kind="continuous", task=None):
    return [None, name, low, high, kind, task]


def task_description(race):
    units = {
        "protoss": "Stalker, Zealot, and Colossus",
        "terran": "Marine, Marauder, and Medivac",
        "zerg": "Zergling, Hydralisk, and Baneling",
    }[race]
    return (
        f"{race}_10_vs_10 is a SMACv2 procedurally generated 10-on-10 battle. "
        f"Both teams are sampled from {units}; starts use surrounded or reflected layouts. "
        "Agents act under partial observability and must communicate role, health, position, "
        "and locally observed enemy information to coordinate focus fire, support, and movement."
    )


def observation_rows(race):
    has_shield = race == "protoss"
    rows = [row(f"move_action_{x}", kind="categorical") for x in ("north", "south", "east", "west")]
    entity_fields = [
        ("shootable", 0, 1, "categorical"),
        ("distance", 0, 1, "continuous"),
        ("relative_x", -1, 1, "continuous"),
        ("relative_y", -1, 1, "continuous"),
        ("health", 0, 1, "continuous"),
    ]
    if has_shield:
        entity_fields.append(("shield", 0, 1, "continuous"))
    entity_fields += [(f"unit_type_bit_{i}", 0, 1, "categorical") for i in range(3)]
    for enemy in range(10):
        rows += [row(f"enemy_{enemy}_{name}", low, high, kind) for name, low, high, kind in entity_fields]
    ally_fields = [("visible", *values) if name == "shootable" else (name, *values)
                   for name, *values in entity_fields]
    for ally in range(9):
        rows += [row(f"ally_{ally}_{name}", low, high, kind) for name, low, high, kind in ally_fields]
    rows.append(row("own_health"))
    if has_shield:
        rows.append(row("own_shield"))
    rows += [row(f"own_unit_type_bit_{i}", kind="categorical") for i in range(3)]
    rows += [row("own_pos_x", -1, 1), row("own_pos_y", -1, 1)]
    actions = ["no-op", "stop", "move_north", "move_south", "move_east", "move_west"]
    actions += [f"attack_enemy_{i}" for i in range(10)]
    rows += [row(f"Last Action - {name}", kind="categorical") for name in actions]
    rows += [row(f"Is this agent{i}", kind="categorical") for i in range(10)]
    rows[0][-1] = task_description(race)
    return rows


def state_rows(race):
    has_shield = race == "protoss"
    rows = []
    ally_fields = [
        ("health", 0, 1, "continuous"), ("cooldown", 0, 1, "continuous"),
        ("absolute_x", -1, 1, "continuous"), ("absolute_y", -1, 1, "continuous"),
    ]
    if has_shield:
        ally_fields.append(("shield", 0, 1, "continuous"))
    ally_fields += [(f"unit_type_bit_{i}", 0, 1, "categorical") for i in range(3)]
    enemy_fields = [
        ("health", 0, 1, "continuous"), ("absolute_x", -1, 1, "continuous"),
        ("absolute_y", -1, 1, "continuous"),
    ]
    if has_shield:
        enemy_fields.append(("shield", 0, 1, "continuous"))
    enemy_fields += [(f"unit_type_bit_{i}", 0, 1, "categorical") for i in range(3)]
    for ally in range(10):
        rows += [row(f"ally_{ally}_{name}", low, high, kind) for name, low, high, kind in ally_fields]
    for enemy in range(10):
        rows += [row(f"enemy_{enemy}_{name}", low, high, kind) for name, low, high, kind in enemy_fields]
    action_names = ["no-op", "stop", "move_north", "move_south", "move_east", "move_west"]
    action_names += [f"attack_enemy_{i}" for i in range(10)]
    for ally in range(10):
        rows += [row(f"ally_{ally}_last_action_{name}", kind="categorical") for name in action_names]
    rows[0][-1] = task_description(race)
    return rows


def write_sheet(workbook, name, rows):
    if name in workbook.sheetnames:
        del workbook[name]
    sheet = workbook.create_sheet(name)
    for index, values in enumerate(rows, start=1):
        values[0] = index
        sheet.append(values)


def main():
    backup = BOOK.with_name("sc2.before_smacv2_10v10.xlsx")
    if not backup.exists():
        copy2(BOOK, backup)
    workbook = load_workbook(BOOK)
    expected = {"protoss": (208, 310), "terran": (188, 290), "zerg": (188, 290)}
    for race, (obs_dim, state_dim) in expected.items():
        obs = observation_rows(race)
        state = state_rows(race)
        assert len(obs) == obs_dim, (race, len(obs), obs_dim)
        assert len(state) == state_dim, (race, len(state), state_dim)
        write_sheet(workbook, f"{race}_10_vs_10", obs)
        write_sheet(workbook, f"{race}_10_vs_10(state)", state)
        print(f"{race}: obs={len(obs)} state={len(state)}")
    workbook.save(BOOK)
    print(f"saved: {BOOK}")


if __name__ == "__main__":
    main()
